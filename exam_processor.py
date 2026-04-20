import copy
import datetime as dt
import io
import json
import re
from collections import Counter, defaultdict
from typing import Dict, List, Optional

import fitz
import openpyxl
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz, process
from unidecode import unidecode

SHEET_NAME = "EXAMES"

PREFERRED_SHEET_LABELS = {
    "CÁLCIO": ["CALCIO IONIZADO", "CALCIO"],
    "CÁLCIO IONIZADO": ["CALCIO IONIZADO"],
    "FERRO": ["FERRO SERICO", "FERRO"],
    "SATURAÇÃO DA TRANSFERRINA": ["INDICE DE SATURACAO DE TRANSFERRINA", "SATURACAO DA TRANSFERRINA"],
}

ROW_PRIORITY = {
    "CÁLCIO IONIZADO": 100,
    "CÁLCIO": 50,
}

IGNORE_IN_SCAN = (
    "METODO", "MATERIAL", "VALOR", "REFER", "NOTA", "COLETA", "LIBERACAO",
    "LIBERADO", "RESPONSAVEL", "CONVENIO", "UNIDADE", "CODIGO", "CPF", "MEDICO",
    "PAG.", "PAGINA", "ASSINATURA", "ENDERECO", "DATA DA GERACAO", "NUCLEO DE ATENDIMENTO",
)


def normalize_text(text: str) -> str:
    text = unidecode(text or "").upper()
    text = text.replace("\r", "")
    text = text.replace("–", "-").replace("—", "-").replace("‑", "-")
    text = text.replace("²", "2").replace("³", "3")
    return text


def normalize_key(text: str) -> str:
    text = normalize_text(text)
    text = re.sub(r"[^A-Z0-9/% ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def convert_value(raw: str):
    raw = normalize_text(str(raw)).strip()
    m = re.search(r"\d[\d.,]*", raw)
    if not m:
        return raw
    num = m.group(0)
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", num):
        num = num.replace(".", "").replace(",", ".")
    elif "," in num and "." not in num:
        num = num.replace(",", ".")
    elif "." in num and "," in num:
        if num.rfind(",") > num.rfind("."):
            num = num.replace(".", "").replace(",", ".")
        else:
            num = num.replace(",", "")
    try:
        value = float(num)
        return int(value) if value.is_integer() else value
    except Exception:
        return num


def extract_pdf_pages(pdf_stream) -> List[str]:
    data = pdf_stream.read() if hasattr(pdf_stream, "read") else pdf_stream
    pages = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            pages.append(page.get_text("text"))
    return pages


def cleaned_lines(page_text: str) -> List[str]:
    lines = []
    for line in normalize_text(page_text).splitlines():
        line = re.sub(r"\s+", " ", line).strip()
        if line:
            lines.append(line)
    return lines


def looks_metadata(line: str) -> bool:
    return any(token in line for token in IGNORE_IN_SCAN)


def extract_from_line(line: str, units: Optional[List[str]] = None) -> Optional[float]:
    if units:
        if not any(unit in line for unit in units):
            return None
    m = re.search(r"(?:[<>]=?\s*)?(?:ACIMA DE|SUPERIOR A|INFERIOR A|MENOR QUE|MAIOR QUE|ATE)?\s*\d[\d.,]*", line)
    if not m:
        return None
    return convert_value(m.group(0))


def line_matches_title(line: str, title: str) -> bool:
    stripped = line.lstrip("*- ")
    if not stripped.startswith(title):
        return False
    line = stripped
    if title == "GLICOSE" and ("MEDIA ESTIMADA" in line or "6 FOSFATO" in line):
        return False
    if title == "HEMOGLOBINA A" and ("GLICADA" in line or "A1C" in line or "A2" in line or "F" in line):
        return False
    if title == "LIPOPROTEINA A" and line.startswith("APOLIPOPROTEINA"):
        return False
    return True


def extract_title_window(lines: List[str], titles: List[str], units: Optional[List[str]] = None, lookahead: int = 16) -> Optional[float]:
    def has_units(s: str) -> bool:
        return units is None or any(unit in s for unit in units)

    def extract_number_anywhere(s: str):
        m = re.search(r"(?:[<>]=?\s*)?(?:ACIMA DE|SUPERIOR A|INFERIOR A|MENOR QUE|MAIOR QUE|ATE)?\s*\d[\d.,]*", s)
        if m:
            return convert_value(m.group(0))
        return None

    for i, line in enumerate(lines):
        if not any(line_matches_title(line, title) for title in titles):
            continue

        same_line = extract_from_line(line, units)
        if same_line is not None and not line.startswith("RESULTADO"):
            return same_line

        same_number = extract_number_anywhere(line)
        if same_number is not None and not line.startswith("RESULTADO"):
            if units and i + 1 < len(lines) and has_units(lines[i + 1]):
                return same_number

        window = lines[i:min(i + lookahead, len(lines))]
        saw_result = False
        for j, w in enumerate(window):
            if j == 0:
                continue
            if "VALOR DE REFERENCIA" in w or ("REFERENCIA" in w and "INTERVALO" in w):
                if saw_result:
                    break
            if looks_metadata(w) and "RESULTADO" not in w and "INDICE" not in w and "HOMA" not in w and "CKD-EPI" not in w and "EGFR" not in w:
                continue

            candidate_context = "RESULTADO" in w or "INDICE" in w or "HOMA" in w or (j <= 3 and not looks_metadata(w))
            if not candidate_context:
                continue

            if "RESULTADO" in w:
                saw_result = True

            value = extract_from_line(w, units)
            if value is not None:
                return value

            raw_num = extract_number_anywhere(w)
            if raw_num is not None:
                if has_units(w):
                    return raw_num
                if j + 1 < len(window) and has_units(window[j + 1]):
                    return raw_num

            if "RESULTADO" in w:
                for k in range(j + 1, min(j + 4, len(window))):
                    if "VALOR DE REFERENCIA" in window[k] or ("REFERENCIA" in window[k] and "INTERVALO" in window[k]):
                        break
                    val = extract_from_line(window[k], units)
                    if val is not None:
                        return val
                    raw_next = extract_number_anywhere(window[k])
                    if raw_next is not None:
                        if has_units(window[k]):
                            return raw_next
                        if k + 1 < len(window) and has_units(window[k + 1]):
                            return raw_next
    return None

def extract_with_regex(page_text: str, pattern: str, group: int = 1) -> Optional[float]:
    m = re.search(pattern, normalize_text(page_text), flags=re.M)
    if m:
        return convert_value(m.group(group))
    return None


def extract_hemogram(page_text: str) -> Dict[str, float]:
    text = normalize_text(page_text)
    results = {}
    patterns = {
        "HEMOGLOBINA": [
            (r"^\s*HEMOGLOBINA\s+(?!GLICADA)(\d[\d.,]*)\s*G/DL", 1),
            (r"^\s*HEMOGLOBINA(?: \(G/DL\))?[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.,]*)", 1),
        ],
        "HEMATÓCRITO": [
            (r"^\s*HEMATOCRITO\s+(\d[\d.,]*)\s*%", 1),
            (r"^\s*HEMATOCRITO(?: \(%\))?[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.,]*)", 1),
        ],
        "VCM": [
            (r"^\s*VCM\s+(\d[\d.,]*)\s*FL", 1),
            (r"^\s*VGM(?: \(FL\))?[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.,]*)", 1),
            (r"^\s*VOLUME CORPUSCULAR MEDIO[^\d\n]*:?\s*(\d[\d.,]*)", 1),
        ],
        "HCM": [
            (r"^\s*HCM\s+(\d[\d.,]*)\s*PG", 1),
            (r"^\s*HBGM(?: \(PG\))?[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.,]*)", 1),
            (r"^\s*HEMOGLOBINA CORPUSCULAR MEDIA[^\d\n]*:?\s*(\d[\d.,]*)", 1),
        ],
        "CHCM": [
            (r"^\s*CHCM\s+(\d[\d.,]*)\s*G/DL", 1),
            (r"^\s*CHBGM(?: \(G/DL\))?[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.,]*)", 1),
            (r"^\s*CONCENTRACAO DE HEMOGLOBINA CORPUSCULAR MEDIA[^\d\n]*:?\s*(\d[\d.,]*)", 1),
        ],
        "RDW": [
            (r"^\s*RDW\s+(\d[\d.,]*)\s*%", 1),
            (r"^\s*RDW(?: \(%\))?[^\d\n]*:?\s*(\d[\d.,]*)", 1),
        ],
        "LEUCÓCITOS TOTAIS": [
            (r"^\s*LEUCOCITOS\s+100\s*%\s*(\d[\d.]*)\s*/", 1),
            (r"^\s*LEUCOCITOS\s+(\d[\d.]*)\s+100\b", 1),
            (r"^\s*LEUCOCITOS(?: \(/MM3\))?[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.]*)", 1),
        ],
        "EOSINÓFILOS": [
            (r"^\s*EOSINOFILOS\s+\d[\d.,]*\s*%\s*(\d[\d.]*)\s*/", 1),
            (r"^\s*EOSINOFILOS\s+\d[\d.,]*\s+(\d[\d.]*)\b", 1),
            (r"^\s*EOSINOFILOS[^\d\n]*:?\s*(?:\n\s*)?\d[\d.,]*\s+(\d[\d.]*)", 1),
        ],
        "MONÓCITOS": [
            (r"^\s*MONOCITOS\s+\d[\d.,]*\s*%\s*(\d[\d.]*)\s*/", 1),
            (r"^\s*MONOCITOS\s+\d[\d.,]*\s+(\d[\d.]*)\b", 1),
            (r"^\s*MONOCITOS[^\d\n]*:?\s*(?:\n\s*)?\d[\d.,]*\s+(\d[\d.]*)", 1),
        ],
        "PLAQUETAS": [
            (r"TOTAL DE PLAQUETAS:\s*(\d[\d.]*)/MM3", 1),
            (r"^\s*PLAQUETAS\s+(\d[\d.]*)\s*/", 1),
            (r"^\s*PLAQUETAS[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.]*)\s*/", 1),
            (r"^\s*PLAQUETAS\s+(\d[\d.]*)\s+X 10", 1),
            (r"CONTAGEM DE\s*\n\s*PLAQUETAS\s*\n\s*(\d[\d.]*)\s*/", 1),
        ],
        "VPM PLAQUETAS": [
            (r"VOLUME PLAQUETARIO MEDIO:\s*(\d[\d.,]*)\s*FL", 1),
            (r"^\s*VPM\s+(\d[\d.,]*)\s*FL", 1),
            (r"^\s*VPM[^\d\n]*:?\s*(?:\n\s*)?(\d[\d.,]*)\s*FL", 1),
            (r"^\s*VMP\s+(\d[\d.,]*)\s*FL", 1),
        ],
    }
    for exam, plist in patterns.items():
        for pattern, group in plist:
            val = extract_with_regex(text, pattern, group)
            if val is not None:
                results[exam] = val
                break
    return results


def extract_special_from_page(page_text: str) -> Dict[str, float]:
    lines = cleaned_lines(page_text)
    joined = "\n".join(lines)
    results: Dict[str, float] = {}

    # Bioquímica básica / endocrino / lipídios
    generic_specs = {
        "GLICEMIA DE JEJUM": (["GLICOSE", "GLICEMIA DE JEJUM"], ["MG/DL"]),
        "INSULINA DE JEJUM": (["INSULINA BASAL", "INSULINA"], ["UUI/ML", "MUI/ML", "UU/ML", "MU/L"]),
        "HOMA IR": (["HOMA-IR", "HOMA - IR", "INDICE DE HOMA - IR", "CALCULO DO INDICE DE HOMA", "CALCULO DO HOMA-IR"], None),
        "HOMA BETA": (["HOMA BETA", "HOMA-BETA", "INDICE DE HOMA - BETA"], None),
        "HEMOGLOBINA GLICADA A1C": (["HEMOGLOBINA GLICADA", "HBA1C"], ["%"]),
        "GLICEMIA MÉDIA ESTIMADA": (["GLICEMIA MEDIA ESTIMADA", "GLICOSE MEDIA ESTIMADA", "GME"], ["MG/DL"]),
        "COLESTEROL TOTAL": (["COLESTEROL TOTAL"], ["MG/DL"]),
        "COLESTEROL HDL": (["HDL - COLESTEROL", "HDL COLESTEROL", "HDL-COLESTEROL", "COLESTEROL HDL"], ["MG/DL"]),
        "COLESTEROL LDL": (["LDL - COLESTEROL", "LDL COLESTEROL", "LDL-COLESTEROL", "COLESTEROL LDL"], ["MG/DL"]),
        "TRIGLICERÍDEOS": (["TRIGLICERIDES", "TRIGLICERIDEOS"], ["MG/DL"]),
        "COLESTEROL VLDL": (["VLDL - COLESTEROL", "VLDL COLESTEROL", "VLDL-COLESTEROL", "COLESTEROL VLDL"], ["MG/DL"]),
        "COLESTEROL NÃO HDL": (["NAO HDL - COLESTEROL", "NAO HDL COLESTEROL", "NAO-HDL-COLESTEROL", "COLESTEROL NAO HDL"], ["MG/DL"]),
        "LIPOPROTEÍNA A": (["LIPOPROTEINA A", 'LIPOPROTEINA "A"', 'LIPOPROTEINA LP(A)'], ["MG/DL", "NMOL/L"]),
        "APOLIPOPROTEINA A1": (["APOLIPOPROTEINA A-1", "APOLIPOPROTEINA A1", 'APOLIPOPROTEINA "A"'], ["MG/DL"]),
        "APOLIPOPROTEÍNA B": (["APOLIPOPROTEINA B", 'APOLIPOPROTEINA "B"'], ["MG/DL"]),
        "RELAÇÃO APO B / APO A": (["RAZAO APO B / APO A1", "RELACAO APO B / APO A"], None),
        "PROTEÍNA C REATIVA ULTRASSENSÍVEL (PCR-US)": (["PROTEINA C REATIVA", "PROTEINA C-REATIVA", "PCR"], ["MG/DL"]),
        "VELOCIDADE DE HEMOSSEDIMENTAÇÃO": (["VELOCIDADE DE HEMOSSEDIMENTACAO", "HEMOSSEDIMENTACAO", "VHS 1A HORA", "VHS"], ["MM", "MM/HORA"]),
        "UREIA": (["UREIA"], ["MG/DL"]),
        "CREATININA": (["CREATININA"], ["MG/DL"]),
        "TFG": (["ESTIMATIVA DA TAXA DE FILTRACAO GLOMERULAR", "ESTIMATIVA DA FILTRACAO GLOMERULAR", "TFGE", "EGFR", "CKD-EPI 2021"], ["ML/MIN"]),
        "CÁLCIO": (["CALCIO, SORO", "CALCIO"], ["MG/DL"]),
        "CÁLCIO IONIZADO": (["CALCIO IONIZADO", "CALCIO IONIZAVEL", "CALCIO IONICO"], ["MG/DL", "MMOL/L"]),
        "PARATORMÔNIO MOLÉCULA INTACTA": (["PARATORMONIO"], ["PG/ML"]),
        "FERRITINA": (["FERRITINA"], ["NG/ML", "MICROG/L"]),
        "FERRO": (["FERRO SERICO", "FERRO"], ["UG/DL", "MCG/DL"]),
        "SATURAÇÃO DA TRANSFERRINA": (["SATURACAO DA TRANSFERRINA", "INDICE DE SATURACAO DE TRANSFERRINA"], ["%"]),
        "PLAQUETAS": (["CONTAGEM DE PLAQUETAS", "PLAQUETAS"], ["/UL", "/MM3", "X 10"]),
        "VPM PLAQUETAS": (["VPM", "VMP", "VOLUME PLAQUETARIO MEDIO"], ["FL"]),
        "HEMOGLOBINA A1": (["HEMOGLOBINA A :", "HEMOGLOBINA A"], ["%"]),
        "HEMOGLOBINA A2": (["HEMOGLOBINA A2"], ["%"]),
        "HEMOGLOBINA F": (["HEMOGLOBINA F"], ["%"]),
        "G6PD": (["GLICOSE 6 FOSFATO DESIDROGENASE", "G6PD"], ["U/G HB"]),
    }

    for exam, (titles, units) in generic_specs.items():
        if exam == "COLESTEROL NÃO HDL" and "NAO HDL" not in joined and "NAO-HDL" not in joined:
            continue
        if exam == "RELAÇÃO APO B / APO A":
            m = re.search(r"RAZAO APO B / APO A1\s*:?\s*(\d[\d.,]*)", joined)
            if m:
                results[exam] = convert_value(m.group(1))
            continue
        if exam == "TFG":
            m = re.search(r"CKD-EPI 2021\s*:?\s*(\d[\d.,]*)", joined)
            if m:
                results[exam] = convert_value(m.group(1))
                continue
            m = re.search(r"(?:ACIMA DE|SUPERIOR A|>)\s*(\d[\d.,]*)\s*ML/MIN", joined)
            if m and any(k in joined for k in titles):
                results[exam] = convert_value(m.group(1))
                continue
            val = extract_title_window(lines, titles, units)
            if val is not None:
                results[exam] = val
            continue
        if exam == "GLICEMIA MÉDIA ESTIMADA":
            m = re.search(r"GLICEMIA MEDIA ESTIMADA\s+DE\s+(\d[\d.,]*)\s*MG/DL", joined)
            if m:
                results[exam] = convert_value(m.group(1))
                continue
        if exam == "HOMA IR":
            m = re.search(r"HOMA\s*-?\s*IR[^\n]*\n\s*(\d[\d.,]*)", joined)
            if not m:
                m = re.search(r"INDICE DE HOMA\s*-\s*IR(?:\n[^\n]*){0,35}?RESULTADO:?\s*\n\s*(\d[\d.,]*)", joined)
            if not m:
                m = re.search(r"CALCULO DO HOMA\s*-?\s*IR\s*:?\s*(\d[\d.,]*)", joined)
            if not m:
                m = re.search(r"CALCULO DO HOMAIR\s*:?\s*(\d[\d.,]*)", joined)
            if m:
                results[exam] = convert_value(m.group(1))
                continue
        if exam == "HOMA BETA":
            m = re.search(r"HOMA\s*-?\s*BETA[^\n]*\n\s*(\d[\d.,]*)", joined)
            if not m:
                m = re.search(r"INDICE DE HOMA\s*-\s*BETA(?:\n[^\n]*){0,35}?INDICE\.*:?\s*\n\s*(\d[\d.,]*)", joined)
            if m:
                results[exam] = convert_value(m.group(1))
                continue
        if exam == "SATURAÇÃO DA TRANSFERRINA":
            m = re.search(r"(?:^|\n)SATURACAO DA TRANSFERRINA[^:\n]*:\s*(\d[\d.,]*)\s*%", joined)
            if not m:
                m = re.search(r"(?:^|\n)INDICE DE SATURACAO DE TRANSFERRINA[^:\n]*:\s*(\d[\d.,]*)\s*%", joined)
            if m:
                results[exam] = convert_value(m.group(1))
                continue
        val = extract_title_window(lines, titles, units)
        if val is not None:
            results[exam] = val

    if any(token in joined for token in ["HEMOGRAMA", "ERITROGRAMA", "SERIE VERMELHA", "HEMOGRAMA COMPLETO"]):
        results.update(extract_hemogram(page_text))
    return results


def parse_exams_from_pages(pages: List[str]):
    extracted: Dict[str, Dict[str, object]] = {}
    for page in pages:
        page_results = extract_special_from_page(page)
        for exam, value in page_results.items():
            if exam not in extracted:
                extracted[exam] = {"value": value, "pattern_used": "page_scan"}

    all_possible = [
        "GLICEMIA DE JEJUM", "INSULINA DE JEJUM", "HOMA IR", "HOMA BETA",
        "HEMOGLOBINA GLICADA A1C", "GLICEMIA MÉDIA ESTIMADA", "COLESTEROL TOTAL",
        "COLESTEROL HDL", "COLESTEROL LDL", "TRIGLICERÍDEOS", "COLESTEROL VLDL",
        "COLESTEROL NÃO HDL", "LIPOPROTEÍNA A", "APOLIPOPROTEINA A1",
        "APOLIPOPROTEÍNA B", "RELAÇÃO APO B / APO A", "PROTEÍNA C REATIVA ULTRASSENSÍVEL (PCR-US)",
        "VELOCIDADE DE HEMOSSEDIMENTAÇÃO", "UREIA", "CREATININA", "TFG", "CÁLCIO",
        "CÁLCIO IONIZADO", "PARATORMÔNIO MOLÉCULA INTACTA", "FERRITINA", "FERRO",
        "SATURAÇÃO DA TRANSFERRINA", "HEMOGLOBINA", "HEMATÓCRITO", "HCM", "VCM",
        "CHCM", "RDW", "LEUCÓCITOS TOTAIS", "EOSINÓFILOS", "MONÓCITOS", "PLAQUETAS",
        "VPM PLAQUETAS", "HEMOGLOBINA A1", "HEMOGLOBINA A2", "HEMOGLOBINA F", "G6PD"
    ]
    not_found = [x for x in all_possible if x not in extracted]
    return extracted, not_found


def extract_collection_date(text: str) -> dt.date:
    normalized = normalize_text(text)
    patterns = [
        r"RECEBIDO/COLETADO\s+EM:\s*(\d{2}/\d{2}/\d{4})",
        r"DATA COLETA/RECEBIMENTO:\s*(\d{2}/\d{2}/\d{4})",
        r"\bCOLETA:\s*(\d{2}/\d{2}/\d{4})",
        r"COLETADO EM:\s*(\d{2}/\d{2}/\d{4})",
        r"ENTRADA\.*:\s*(\d{2}/\d{2}/\d{4})",
        r"DATA DA FICHA\s*(\d{2}/\d{2}/\d{4})",
    ]
    matches = []
    for pattern in patterns:
        matches.extend(re.findall(pattern, normalized))
    if matches:
        chosen = Counter(matches).most_common(1)[0][0]
        return dt.datetime.strptime(chosen, "%d/%m/%Y").date()
    return dt.date.today()


def build_row_map(ws):
    row_map = defaultdict(list)
    available_labels = []
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 2).value
        if label:
            key = normalize_key(str(label))
            row_map[key].append(row)
            available_labels.append((key, row, str(label)))
    return row_map, available_labels


def find_row_for_exam(canonical_name: str, row_map, available_labels):
    preferred = PREFERRED_SHEET_LABELS.get(canonical_name, [canonical_name])
    for label in preferred:
        key = normalize_key(label)
        if key in row_map and row_map[key]:
            return row_map[key][0], 100, label

    key = normalize_key(canonical_name)
    if key in row_map and row_map[key]:
        return row_map[key][0], 100, canonical_name

    choices = {}
    for label_key, row, original in available_labels:
        if label_key not in choices:
            choices[label_key] = (row, original)

    best = process.extractOne(key, list(choices.keys()), scorer=fuzz.ratio)
    if best and best[1] >= 86:
        row, original = choices[best[0]]
        return row, best[1], original
    return None, 0, None


def find_data_row(ws) -> int:
    for row in range(1, ws.max_row + 1):
        if normalize_key(str(ws.cell(row, 2).value or "")) == "DATA":
            return row
    raise ValueError("A linha 'DATA' não foi encontrada na coluna B da planilha.")


def next_available_data_column(ws, data_row: int, start_col: int = 3) -> int:
    last_used = start_col - 1
    for col in range(start_col, ws.max_column + 1):
        if ws.cell(data_row, col).value not in (None, ""):
            last_used = col
    return max(start_col, last_used + 1)


def clone_column_format(ws, source_col: int, target_col: int):
    src_letter = get_column_letter(source_col)
    dst_letter = get_column_letter(target_col)
    ws.column_dimensions[dst_letter].width = ws.column_dimensions[src_letter].width
    ws.column_dimensions[dst_letter].hidden = ws.column_dimensions[src_letter].hidden
    for row in range(1, ws.max_row + 1):
        src = ws.cell(row, source_col)
        dst = ws.cell(row, target_col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.font:
            dst.font = copy.copy(src.font)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.border:
            dst.border = copy.copy(src.border)
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.protection:
            dst.protection = copy.copy(src.protection)
        dst.number_format = src.number_format


def write_to_workbook(xlsx_stream, extracted_data: dict, collection_date: dt.date):
    wb = openpyxl.load_workbook(xlsx_stream)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"A planilha precisa conter a aba '{SHEET_NAME}'.")
    ws = wb[SHEET_NAME]

    row_map, available_labels = build_row_map(ws)
    data_row = find_data_row(ws)
    target_col = next_available_data_column(ws, data_row)
    reference_col = max(3, target_col - 1)
    clone_column_format(ws, reference_col, target_col)

    ws.cell(data_row, target_col).value = collection_date
    ws.cell(data_row, target_col).number_format = "dd/mm/yyyy"

    written = []
    unmatched = []
    occupied_rows = {}

    for exam, payload in extracted_data.items():
        row, score, matched_label = find_row_for_exam(exam, row_map, available_labels)
        if row is None:
            unmatched.append({"exam": exam, "value": payload["value"]})
            continue

        existing = occupied_rows.get(row)
        priority = ROW_PRIORITY.get(exam, 10)
        if existing and ROW_PRIORITY.get(existing["exam"], 10) >= priority:
            continue

        ws.cell(row, target_col).value = payload["value"]
        occupied_rows[row] = {"exam": exam, "value": payload["value"]}

        existing_item = next((item for item in written if item["row"] == row), None)
        info = {
            "exam": exam,
            "value": payload["value"],
            "row": row,
            "column": get_column_letter(target_col),
            "match_score": score,
            "sheet_match": matched_label,
        }
        if existing_item:
            existing_item.update(info)
        else:
            written.append(info)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, {
        "target_column": get_column_letter(target_col),
        "data_row": data_row,
        "collection_date": collection_date.strftime("%d/%m/%Y"),
        "written_to_excel": written,
        "found_in_pdf_but_not_matched_in_sheet": unmatched,
    }


def process_files(pdf_stream, xlsx_stream, pdf_filename: str = "exames.pdf", xlsx_filename: str = "painel.xlsx"):
    pages = extract_pdf_pages(pdf_stream)
    extracted, not_found = parse_exams_from_pages(pages)
    collection_date = extract_collection_date("\n".join(pages))
    workbook_stream, write_info = write_to_workbook(xlsx_stream, extracted, collection_date)
    report = {
        "pdf": pdf_filename,
        "xlsx_template": xlsx_filename,
        "target_column": write_info["target_column"],
        "collection_date": write_info["collection_date"],
        "total_exams_extracted": len(extracted),
        "written_to_excel": write_info["written_to_excel"],
        "not_found_in_pdf": not_found,
        "found_in_pdf_but_not_matched_in_sheet": write_info["found_in_pdf_but_not_matched_in_sheet"],
    }
    return workbook_stream, report


def report_to_json_bytes(report: dict) -> bytes:
    return json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8")
